# -*- coding: utf-8 -*-
"""Genera la Especificación de Requerimientos Funcionales de LigaPlus en .docx."""
import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VERDE = RGBColor(0x10, 0x3D, 0x2B)
VERDE2 = RGBColor(0x1F, 0x7A, 0x4D)
NARANJA = RGBColor(0xC4, 0x5F, 0x1E)
GRIS = RGBColor(0x44, 0x44, 0x41)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FECHA = "18 de junio de 2026"

doc = Document()


# ------------------------------------------------------------------ helpers
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=10, align=None, white=False):
    cell.text = ''
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    if white:
        run.font.color.rgb = BLANCO
    elif color is not None:
        run.font.color.rgb = color


def heading_border(style, hex_color):
    pPr = style.element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '10')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), hex_color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_field(paragraph, instr, placeholder=''):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = placeholder
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2); run._r.append(t); run._r.append(f3)


def enable_update_fields():
    el = doc.settings.element
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    el.append(uf)


# ------------------------------------------------------------------ styles
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0x20, 0x20, 0x1E)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for name, color, size in [('Heading 1', VERDE, 17), ('Heading 2', VERDE2, 13.5), ('Heading 3', NARANJA, 11.5)]:
    st = styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = color
    st.font.size = Pt(size)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14 if name == 'Heading 1' else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
heading_border(styles['Heading 1'], '1F7A4D')

ti = styles['Title']
ti.font.name = 'Calibri'
ti.font.size = Pt(34)
ti.font.color.rgb = VERDE
ti.font.bold = True

sec = doc.sections[0]
sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.3); sec.right_margin = Cm(2.3)


# ------------------------------------------------------------------ portada
band = doc.add_table(rows=1, cols=1)
band.alignment = WD_TABLE_ALIGNMENT.CENTER
band.columns[0].width = Cm(16.4)
c = band.rows[0].cells[0]
shade(c, '103D2B')
c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
r = c.paragraphs[0].add_run('LigaPlus')
r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = BLANCO; r.font.name = 'Calibri'
sub = c.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run('Plataforma SaaS para la gestión integral de ligas deportivas')
rs.font.size = Pt(11); rs.font.color.rgb = RGBColor(0xC6, 0xDC, 0x4D); rs.font.name = 'Calibri'
c.paragraphs[0].paragraph_format.space_before = Pt(10)
sub.paragraph_format.space_after = Pt(12)

doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
rt = t.add_run('Especificación de Requerimientos Funcionales')
rt.font.size = Pt(24); rt.font.bold = True; rt.font.color.rgb = VERDE; rt.font.name = 'Calibri'
st2 = doc.add_paragraph(); st2.alignment = WD_ALIGN_PARAGRAPH.CENTER
rst = st2.add_run('Descripción detallada de cada funcionalidad del sistema')
rst.font.size = Pt(12); rst.font.color.rgb = GRIS; rst.italic = True

doc.add_paragraph(); doc.add_paragraph()
meta = doc.add_table(rows=6, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = 'Light List Accent 1'
datos = [
    ('Proyecto', 'LigaPlus — Gestión de ligas deportivas amateur'),
    ('Versión del documento', '1.0'),
    ('Fecha', FECHA),
    ('Estado', 'Borrador para revisión'),
    ('Preparado por', 'Equipo LigaPlus'),
    ('Clasificación', 'Confidencial — uso interno y comité'),
]
for i, (k, v) in enumerate(datos):
    set_cell_text(meta.rows[i].cells[0], k, bold=True, color=VERDE, size=10.5)
    set_cell_text(meta.rows[i].cells[1], v, size=10.5)
    meta.columns[0].width = Cm(5.0); meta.columns[1].width = Cm(11.0)

doc.add_page_break()


# ------------------------------------------------------------------ control de versiones
doc.add_heading('Control de versiones', level=1)
cv = doc.add_table(rows=2, cols=4)
cv.style = 'Light Grid Accent 1'
for j, h in enumerate(['Versión', 'Fecha', 'Descripción del cambio', 'Autor']):
    set_cell_text(cv.rows[0].cells[j], h, bold=True, white=True, size=10)
    shade(cv.rows[0].cells[j], '103D2B')
fila = ['1.0', FECHA, 'Versión inicial del documento de requerimientos.', 'Equipo LigaPlus']
for j, v in enumerate(fila):
    set_cell_text(cv.rows[1].cells[j], v, size=10)
cv.columns[0].width = Cm(2.0); cv.columns[1].width = Cm(3.4)
cv.columns[2].width = Cm(8.6); cv.columns[3].width = Cm(3.0)

doc.add_paragraph()
doc.add_heading('Índice', level=1)
intro_toc = doc.add_paragraph()
ri = intro_toc.add_run('Para actualizar el índice: clic derecho sobre él → Actualizar campos → Actualizar toda la tabla (o tecla F9).')
ri.italic = True; ri.font.size = Pt(9); ri.font.color.rgb = GRIS
toc_p = doc.add_paragraph()
add_field(toc_p, 'TOC \\o "1-2" \\h \\z \\u', 'El índice se generará al abrir el documento en Word.')
doc.add_page_break()


# ------------------------------------------------------------------ render RF
def render_req(rf):
    doc.add_heading(f"{rf['id']} — {rf['nombre']}", level=3)
    doc.add_paragraph(rf['desc'])
    pa = doc.add_paragraph()
    ra = pa.add_run('Actores: '); ra.bold = True; ra.font.color.rgb = VERDE2
    pa.add_run(rf['actores']).italic = True
    if rf.get('reglas'):
        pr = doc.add_paragraph(); prr = pr.add_run('Reglas de negocio'); prr.bold = True; prr.font.color.rgb = VERDE2
        for b in rf['reglas']:
            doc.add_paragraph(b, style='List Bullet')
    if rf.get('aceptacion'):
        pc = doc.add_paragraph(); pcc = pc.add_run('Criterios de aceptación'); pcc.bold = True; pcc.font.color.rgb = VERDE2
        for b in rf['aceptacion']:
            doc.add_paragraph(b, style='List Bullet')


# ================================================================== 1. INTRODUCCIÓN
doc.add_heading('1. Introducción', level=1)
doc.add_heading('1.1 Propósito', level=2)
doc.add_paragraph(
    'Este documento describe en detalle las funcionalidades del sistema LigaPlus, una plataforma SaaS '
    'multi-tenant para la gestión integral de ligas deportivas amateur. Su objetivo es servir como referencia '
    'funcional para el equipo de producto, desarrollo, control de calidad y para la presentación del alcance del '
    'sistema ante terceros (clientes, socios e inversionistas).')
doc.add_heading('1.2 Alcance', level=2)
doc.add_paragraph(
    'LigaPlus cubre el ciclo completo de operación de una liga: configuración de torneos y calendario, '
    'gestión de clubes y jugadores, actas y resultados en vivo, disciplina, designación y pago de personal, '
    'cobros a clubes, comunicación con la comunidad (hinchas) y reportería. El sistema soporta múltiples ligas '
    'de forma aislada sobre una misma infraestructura.')
doc.add_heading('1.3 Definiciones y abreviaturas', level=2)
doc.add_paragraph('El glosario completo se encuentra en la sección 6. Algunas siglas usadas: '
                  'RF (Requerimiento Funcional), RNF (Requerimiento No Funcional), RLS (Row-Level Security), '
                  'PWA (Progressive Web App), JWT (JSON Web Token).')
doc.add_heading('1.4 Convenciones', level=2)
doc.add_paragraph(
    'Cada funcionalidad se identifica con un código RF-AREA-NN. Cada requerimiento incluye una descripción, '
    'los actores que lo utilizan, sus reglas de negocio y los criterios de aceptación que permiten verificar '
    'su correcto funcionamiento.')

doc.add_page_break()

# ================================================================== 2. DESCRIPCIÓN GENERAL
doc.add_heading('2. Descripción general del producto', level=1)
doc.add_heading('2.1 Perspectiva del producto', level=2)
doc.add_paragraph(
    'LigaPlus es una aplicación web responsiva (con capacidades de aplicación instalable y uso sin conexión) '
    'compuesta por un backend de servicios y un frontend para los distintos roles. Cada liga opera como un '
    'inquilino (tenant) independiente: sus datos están aislados a nivel del motor de base de datos, de modo que '
    'una liga nunca accede a información de otra.')
doc.add_heading('2.2 Roles de usuario', level=2)
doc.add_paragraph('El sistema reconoce seis roles principales:')
for it in [
    'Super Administrador (equipo LigaPlus): administra todas las ligas, planes y facturación de la plataforma.',
    'Administrador de liga: opera su liga (torneos, calendario, designaciones, finanzas, disciplina).',
    'Delegado / Club: gestiona el plantel de su club y sus pagos.',
    'Árbitro / Personal: consulta sus designaciones, confirma asistencia y revisa sus pagos.',
    'Jugador: consulta su ficha, calendario y estadísticas.',
    'Hincha: sigue a su club en el portal público de la liga.',
]:
    doc.add_paragraph(it, style='List Bullet')
doc.add_heading('2.3 Supuestos y restricciones', level=2)
for it in [
    'El sistema opera en modalidad multi-tenant; toda funcionalidad respeta el aislamiento entre ligas.',
    'El acceso requiere autenticación, salvo las vistas públicas destinadas a los hinchas.',
    'La plataforma está orientada a uso móvil en cancha y de escritorio en oficina.',
    'Los pagos en línea dependen de la pasarela contratada por cada liga y de la normativa tributaria vigente (Chile).',
]:
    doc.add_paragraph(it, style='List Bullet')

doc.add_page_break()

# ================================================================== 3. REQUERIMIENTOS FUNCIONALES
doc.add_heading('3. Requerimientos funcionales', level=1)
doc.add_paragraph(
    'Esta sección detalla cada funcionalidad del sistema agrupada por módulo. El número del módulo (3.x) '
    'corresponde a un área funcional; los requerimientos individuales se identifican con su código RF.')

MODULOS = [
    ("3.1 Autenticación y gestión de cuentas",
     "Controla el acceso de todos los usuarios al sistema y la seguridad de sus credenciales.",
     [
        dict(id="RF-AUTH-01", nombre="Inicio de sesión",
             desc="Permite a un usuario autenticarse con su correo y contraseña para obtener acceso a su portal según su rol.",
             actores="Todos los usuarios autenticados",
             reglas=["La sesión se sostiene con un token de acceso de corta duración (15 minutos) y un token de refresco de 7 días que rota en cada uso.",
                     "El intento de inicio de sesión está limitado a 5 intentos por IP cada 15 minutos para mitigar ataques de fuerza bruta.",
                     "Las contraseñas se almacenan con hash robusto (bcrypt); nunca en texto plano."],
             aceptacion=["Con credenciales válidas el usuario accede a su portal correspondiente.",
                         "Tras superar el límite de intentos, el sistema bloquea temporalmente nuevos intentos desde esa IP."]),
        dict(id="RF-AUTH-02", nombre="Recuperación de contraseña",
             desc="Permite al usuario restablecer su contraseña cuando la ha olvidado, mediante un enlace enviado a su correo.",
             actores="Todos los usuarios",
             reglas=["El enlace de restablecimiento contiene un token de un solo uso con expiración.",
                     "Al restablecer, la nueva contraseña debe cumplir la política de seguridad (RF-AUTH-03).",
                     "Por privacidad, el sistema no revela si un correo está o no registrado."],
             aceptacion=["El usuario recibe el correo y puede definir una nueva contraseña válida.",
                         "Un enlace usado o expirado es rechazado."]),
        dict(id="RF-AUTH-03", nombre="Política de contraseñas seguras",
             desc="Garantiza que todas las contraseñas tengan un nivel de seguridad acorde a un sistema productivo, siguiendo lineamientos NIST.",
             actores="Todos los usuarios que definen o cambian contraseña",
             reglas=["Longitud mínima de 10 y máxima de 128 caracteres.",
                     "Se rechazan contraseñas comunes (lista negra), secuencias triviales y contraseñas con menos de 4 caracteres distintos.",
                     "Se rechaza que la contraseña contenga el correo o el nombre del usuario.",
                     "El frontend muestra un medidor de fortaleza en tiempo real."],
             aceptacion=["Una contraseña que no cumple la política es rechazada con un mensaje claro.",
                         "El medidor refleja la fortaleza mientras el usuario escribe."]),
        dict(id="RF-AUTH-04", nombre="Invitación y activación de usuarios",
             desc="Permite incorporar delegados, personal y miembros del staff mediante un enlace de invitación con el que definen su contraseña y activan su cuenta.",
             actores="Administrador de liga, Super Administrador",
             reglas=["La invitación se envía por correo y vincula al usuario con su rol y liga.",
                     "La contraseña definida en la activación cumple la política de seguridad.",
                     "El correo del invitado debe ser único dentro de la liga."],
             aceptacion=["El invitado activa su cuenta y accede con el rol asignado.",
                         "No es posible activar dos cuentas con el mismo correo en la misma liga."]),
        dict(id="RF-AUTH-05", nombre="Persistencia y cierre de sesión",
             desc="Mantiene la sesión activa de forma segura entre recargas y permite cerrarla revocando los tokens.",
             actores="Todos los usuarios autenticados",
             reglas=["Al recargar la aplicación la sesión se conserva mientras el token de refresco sea válido.",
                     "El cierre de sesión invalida el token de refresco asociado."],
             aceptacion=["Recargar la página no expulsa al usuario con sesión vigente.",
                         "Tras cerrar sesión, el token revocado no permite obtener nuevos accesos."]),
     ]),

    ("3.2 Administración de la plataforma (Super Administrador)",
     "Funciones exclusivas del equipo LigaPlus para administrar el conjunto de ligas y el negocio de la plataforma.",
     [
        dict(id="RF-PLT-01", nombre="Gestión de ligas (tenants)",
             desc="Permite crear, configurar, activar o suspender ligas, cada una con su propio dominio y datos aislados.",
             actores="Super Administrador",
             reglas=["Cada liga se identifica por su dominio/host y mantiene sus datos completamente aislados de las demás.",
                     "La suspensión de una liga inhabilita el acceso de sus usuarios sin eliminar su información."],
             aceptacion=["El alta de una liga deja operativa su instancia con su administrador inicial.",
                         "Una liga suspendida no permite operación hasta su reactivación."]),
        dict(id="RF-PLT-02", nombre="Planes de suscripción y facturación a ligas",
             desc="Administra los planes que contratan las ligas y la facturación recurrente de la plataforma hacia ellas.",
             actores="Super Administrador",
             reglas=["Cada liga tiene asociado un plan que define su nivel de servicio.",
                     "El sistema registra el estado de pago/facturación de cada liga."],
             aceptacion=["Es posible asignar y cambiar el plan de una liga.",
                         "El estado de facturación de cada liga es consultable."]),
        dict(id="RF-PLT-03", nombre="Impersonación de soporte",
             desc="Permite al Super Administrador operar temporalmente como administrador de una liga para brindar soporte, dejando registro de la sesión.",
             actores="Super Administrador",
             reglas=["Toda sesión de impersonación queda registrada en la auditoría.",
                     "Durante la impersonación se bloquean acciones sensibles (cambio de contraseña, eliminación de cuenta)."],
             aceptacion=["El soporte puede reproducir la vista del administrador de la liga.",
                         "Las acciones sensibles permanecen bloqueadas e identificadas como impersonación."]),
        dict(id="RF-PLT-04", nombre="Estado y versión del sistema",
             desc="Expone el estado de salud y la versión desplegada para verificar la operación y los despliegues.",
             actores="Super Administrador, Operaciones",
             reglas=["Existen verificaciones de vida y de disponibilidad de dependencias.",
                     "Se expone la versión (identificador de despliegue) en ejecución."],
             aceptacion=["Las verificaciones de salud responden el estado real del servicio.",
                         "La versión desplegada es consultable."]),
     ]),

    ("3.3 Configuración de la liga",
     "Parámetros que el administrador define para adecuar la plataforma a su liga.",
     [
        dict(id="RF-LIGA-01", nombre="Temporadas",
             desc="Permite organizar la competencia en temporadas, dentro de las cuales se crean los torneos.",
             actores="Administrador de liga",
             reglas=["Los torneos pertenecen a una temporada.",
                     "Una liga puede tener temporadas históricas y una vigente."],
             aceptacion=["Es posible crear una temporada y asociarle torneos."]),
        dict(id="RF-LIGA-02", nombre="Identidad visual (branding)",
             desc="Permite configurar el logo y los colores de la liga, incluyendo un selector de color con paleta nativa.",
             actores="Administrador de liga",
             reglas=["Los colores definidos se reflejan en los portales de la liga.",
                     "El selector ofrece la paleta de color del sistema operativo del usuario."],
             aceptacion=["Al guardar, la identidad visual se aplica en la liga.",
                         "El cambio de color persiste tras recargar."]),
        dict(id="RF-LIGA-03", nombre="Recintos y canchas",
             desc="Administra los recintos y canchas disponibles para programar los partidos, incluyendo su disponibilidad.",
             actores="Administrador de liga",
             reglas=["Cada cancha tiene un estado de disponibilidad que condiciona la programación.",
                     "El cambio de estado de una cancha se persiste de forma confiable."],
             aceptacion=["Marcar una cancha como disponible se mantiene tras recargar la vista.",
                         "Las canchas no disponibles no se ofrecen para programar."]),
        dict(id="RF-LIGA-04", nombre="Días no jugables y feriados",
             desc="Permite definir fechas en las que no se programan partidos (feriados o días bloqueados).",
             actores="Administrador de liga",
             reglas=["El generador de fixture respeta los días no jugables.",
                     "Los feriados pueden configurarse por liga."],
             aceptacion=["El fixture no asigna partidos en días marcados como no jugables."]),
        dict(id="RF-LIGA-05", nombre="Categorías y series",
             desc="Define las categorías (por edad/nivel) y series en que compiten los equipos.",
             actores="Administrador de liga",
             reglas=["Un torneo puede incluir varias categorías/series con cupos.",
                     "Las inscripciones se asocian a una categoría/serie."],
             aceptacion=["Es posible configurar categorías y series y usarlas al inscribir equipos."]),
     ]),

    ("3.4 Clubes y delegados",
     "Gestión de los clubes participantes y de los delegados responsables de cada uno.",
     [
        dict(id="RF-CLUB-01", nombre="Gestión de clubes",
             desc="Permite registrar y administrar los clubes que participan en la liga.",
             actores="Administrador de liga",
             reglas=["Cada club pertenece a una liga.",
                     "Un club puede inscribir equipos en distintos torneos/categorías."],
             aceptacion=["Es posible crear, editar y listar los clubes de la liga."]),
        dict(id="RF-CLUB-02", nombre="Invitación de delegados",
             desc="Incorpora a los delegados responsables de cada club mediante invitación (ver RF-AUTH-04).",
             actores="Administrador de liga",
             reglas=["El delegado queda vinculado a su club.",
                     "Un delegado gestiona únicamente la información de su club."],
             aceptacion=["El delegado invitado accede al portal de su club."]),
        dict(id="RF-CLUB-03", nombre="Portal del delegado",
             desc="Espacio del delegado para gestionar el plantel de su club, revisar designaciones y consultar sus pagos/cuotas.",
             actores="Delegado / Club",
             reglas=["El delegado solo ve datos de su club.",
                     "Puede consultar el estado de sus cuotas y pagos."],
             aceptacion=["El delegado administra su plantel y consulta su situación financiera."]),
     ]),

    ("3.5 Jugadores e inscripciones",
     "Registro de jugadores, inscripción de equipos a torneos y conformación de planillas.",
     [
        dict(id="RF-JUG-01", nombre="Registro y ficha de jugadores",
             desc="Permite registrar jugadores y mantener su ficha (datos personales y deportivos).",
             actores="Administrador de liga, Delegado",
             reglas=["Cada jugador pertenece a la liga y puede asociarse a un club.",
                     "La ficha respalda las estadísticas y el historial del jugador."],
             aceptacion=["Es posible registrar y consultar la ficha de un jugador."]),
        dict(id="RF-JUG-02", nombre="Inscripción de equipos a torneos",
             desc="Inscribe equipos de un club en un torneo, respetando los cupos por categoría/serie.",
             actores="Administrador de liga, Delegado",
             reglas=["La inscripción respeta el cupo de equipos definido por categoría/serie.",
                     "Un equipo inscrito participa del fixture del torneo."],
             aceptacion=["No es posible exceder el cupo de la categoría/serie."]),
        dict(id="RF-JUG-03", nombre="Planilla del partido",
             desc="Conforma la lista de jugadores habilitados de cada equipo para un partido.",
             actores="Administrador de liga, Delegado, Personal de planilla",
             reglas=["El torneo puede exigir un mínimo de jugadores en planilla para iniciar.",
                     "Solo los jugadores en planilla pueden registrar incidencias en el acta."],
             aceptacion=["Un partido no inicia si no se cumple el mínimo de jugadores configurado.",
                         "No se puede atribuir una incidencia a un jugador fuera de la planilla."]),
        dict(id="RF-JUG-04", nombre="Refuerzos",
             desc="Permite incorporar jugadores de refuerzo cuando el torneo lo habilita, hasta una fecha límite.",
             actores="Administrador de liga, Delegado",
             reglas=["Los refuerzos se habilitan por torneo.",
                     "Existe una fecha límite a partir de la cual no se admiten refuerzos."],
             aceptacion=["No se admiten refuerzos pasada la fecha límite configurada."]),
        dict(id="RF-JUG-05", nombre="Tope de jugadores por equipo",
             desc="Limita la cantidad máxima de jugadores inscritos por equipo en un torneo.",
             actores="Administrador de liga",
             reglas=["El tope se define por torneo.",
                     "No se permite inscribir jugadores por encima del tope."],
             aceptacion=["El sistema impide superar el tope configurado."]),
        dict(id="RF-JUG-06", nombre="Jugadores vetados",
             desc="Mantiene una lista de jugadores inhabilitados que bloquea su inscripción o inclusión en planilla.",
             actores="Administrador de liga",
             reglas=["Un jugador vetado no puede inscribirse ni integrar una planilla.",
                     "El veto se informa al intentar incluir al jugador."],
             aceptacion=["El sistema bloquea la inclusión de un jugador vetado y lo informa."]),
        dict(id="RF-JUG-07", nombre="Importación de planteles (CSV)",
             desc="Permite cargar planteles de forma masiva desde un archivo, agilizando el alta de jugadores.",
             actores="Administrador de liga, Delegado",
             reglas=["La importación valida el formato y reporta errores por fila.",
                     "Los jugadores importados quedan asociados al club/equipo indicado."],
             aceptacion=["Un archivo válido crea los jugadores correspondientes; las filas con error se informan."]),
     ]),

    ("3.6 Torneos",
     "Creación y parametrización de los torneos de la liga.",
     [
        dict(id="RF-TOR-01", nombre="Creación y formato del torneo",
             desc="Crea un torneo definiendo su formato, número de ruedas y los puntos por victoria, empate y derrota.",
             actores="Administrador de liga",
             reglas=["Se admiten formatos de liga (todos contra todos) con una o dos ruedas.",
                     "Los puntos por resultado son configurables."],
             aceptacion=["El torneo creado refleja el formato y los puntajes configurados."]),
        dict(id="RF-TOR-02", nombre="Multi-categoría y series con cupo",
             desc="Permite que un torneo agrupe varias categorías y series, cada una con su cupo de equipos.",
             actores="Administrador de liga",
             reglas=["Cada combinación categoría/serie define un cupo.",
                     "Las inscripciones y tablas se calculan por categoría/serie."],
             aceptacion=["El torneo gestiona correctamente sus categorías/series y cupos."]),
        dict(id="RF-TOR-03", nombre="Criterios de desempate configurables",
             desc="Define el orden de los criterios de desempate de la tabla de posiciones.",
             actores="Administrador de liga",
             reglas=["Los criterios disponibles incluyen puntos, diferencia de gol, goles a favor/en contra, partidos ganados y enfrentamiento directo.",
                     "El orden de los criterios es configurable por torneo."],
             aceptacion=["La tabla ordena los equipos según los criterios configurados."]),
        dict(id="RF-TOR-04", nombre="Parámetros del partido",
             desc="Configura la duración de los periodos y del entretiempo para los partidos del torneo.",
             actores="Administrador de liga",
             reglas=["La duración de periodo y entretiempo se define por torneo.",
                     "Estos valores alimentan el cronómetro del Match Center."],
             aceptacion=["El cronómetro del partido usa la duración configurada en el torneo."]),
        dict(id="RF-TOR-05", nombre="Regla de suspensión por amarillas",
             desc="Define cuántas tarjetas amarillas acumuladas en el torneo generan una fecha de suspensión.",
             actores="Administrador de liga",
             reglas=["El umbral de amarillas para suspensión es configurable por torneo.",
                     "El tribunal aplica la suspensión automáticamente al alcanzarse el umbral."],
             aceptacion=["Al alcanzar el umbral configurado, el jugador queda suspendido la cantidad de fechas correspondiente."]),
        dict(id="RF-TOR-06", nombre="Cobertura de personal por jornada",
             desc="Define cuántos paramédicos y personal adicional debe designar la auto-asignación por jornada.",
             actores="Administrador de liga",
             reglas=["La cantidad de paramédicos y de personal adicional por jornada es configurable.",
                     "Estos parámetros guían la auto-asignación de cobertura de recinto (RF-DES-03)."],
             aceptacion=["La auto-asignación designa la cantidad de personal configurada por jornada."]),
        dict(id="RF-TOR-07", nombre="Estados del torneo",
             desc="Gestiona el ciclo de vida del torneo: borrador, activo y cerrado.",
             actores="Administrador de liga",
             reglas=["En borrador se configura sin afectar la operación.",
                     "Al activarse, el torneo opera con fixture y actas; al cerrarse, se consolida."],
             aceptacion=["El torneo transita correctamente entre sus estados."]),
     ]),

    ("3.7 Fixture y calendario",
     "Generación y mantenimiento del calendario de partidos.",
     [
        dict(id="RF-FIX-01", nombre="Generación automática del fixture",
             desc="Genera automáticamente el calendario de partidos del torneo según su formato y número de ruedas.",
             actores="Administrador de liga",
             reglas=["El generador produce los enfrentamientos para una o dos ruedas.",
                     "Respeta categorías/series del torneo."],
             aceptacion=["El fixture generado cubre todos los enfrentamientos del formato elegido."]),
        dict(id="RF-FIX-02", nombre="Asignación de horarios y canchas",
             desc="Asigna fecha, hora y cancha a cada partido respetando la disponibilidad y los días no jugables.",
             actores="Administrador de liga",
             reglas=["No se programan partidos en días no jugables/feriados.",
                     "Se respeta la disponibilidad de canchas."],
             aceptacion=["Ningún partido queda programado en un día bloqueado o en una cancha no disponible."]),
        dict(id="RF-FIX-03", nombre="Reprogramación y suspensión",
             desc="Permite reprogramar o suspender partidos y fechas, incluyendo reordenamiento por arrastrar y soltar.",
             actores="Administrador de liga",
             reglas=["La reprogramación conserva el historial del partido.",
                     "La suspensión de una fecha afecta a los partidos involucrados."],
             aceptacion=["Un partido reprogramado refleja su nueva fecha sin perder su información."]),
        dict(id="RF-FIX-04", nombre="Walkover",
             desc="Permite declarar un partido por no presentación (walkover), asignando el resultado reglamentario.",
             actores="Administrador de liga",
             reglas=["El walkover asigna directamente el resultado reglamentario (3-0) sin requerir incidencias.",
                     "Es el único caso en que se registran goles sin un goleador asociado."],
             aceptacion=["Un walkover queda con el marcador reglamentario y se distingue de un cierre normal de acta."]),
     ]),

    ("3.8 Acta del partido",
     "Registro oficial del desarrollo y resultado de cada partido.",
     [
        dict(id="RF-ACTA-01", nombre="Certificación de planteles presentes",
             desc="Registra qué jugadores de cada equipo están presentes y habilitados antes de iniciar el partido.",
             actores="Personal de planilla, Administrador de liga",
             reglas=["Solo jugadores presentes y en planilla pueden participar en incidencias.",
                     "Se valida el mínimo de jugadores para iniciar (RF-JUG-03)."],
             aceptacion=["El acta refleja los planteles certificados de ambos equipos."]),
        dict(id="RF-ACTA-02", nombre="Registro de incidencias",
             desc="Registra los hechos del partido: goles, autogoles y tarjetas amarillas y rojas.",
             actores="Personal de planilla, Administrador de liga",
             reglas=["Todo gol o autogol debe tener asignado el jugador que lo marcó; no se puede asignar un gol a un equipo sin el goleador.",
                     "La única excepción al goleador obligatorio es el walkover (RF-FIX-04)."],
             aceptacion=["Intentar registrar un gol sin jugador es rechazado con un mensaje claro.",
                         "Las incidencias quedan asociadas al jugador y equipo correctos."]),
        dict(id="RF-ACTA-03", nombre="Marcador derivado de incidencias",
             desc="Calcula el marcador del partido automáticamente a partir de las incidencias de gol registradas.",
             actores="Sistema",
             reglas=["El marcador se recalcula al agregar o quitar una incidencia de gol.",
                     "El marcador en vivo se difunde al Match Center (RF-MC-02)."],
             aceptacion=["Al registrar o eliminar un gol, el marcador se actualiza de inmediato en todas las vistas."]),
        dict(id="RF-ACTA-04", nombre="Atribución de jugador a una incidencia",
             desc="Permite corregir o asignar el jugador responsable de una incidencia mientras el acta esté abierta.",
             actores="Personal de planilla, Administrador de liga",
             reglas=["El jugador atribuido debe pertenecer a la planilla del equipo correspondiente.",
                     "No se permite la atribución si el acta está cerrada."],
             aceptacion=["La incidencia queda asociada al jugador correcto.",
                         "No es posible modificar incidencias de un acta cerrada."]),
        dict(id="RF-ACTA-05", nombre="Cierre del acta y sanciones",
             desc="Cierra el acta del partido, consolidando el resultado y aplicando automáticamente las sanciones disciplinarias.",
             actores="Administrador de liga",
             reglas=["No se permite cerrar el acta si existen goles sin goleador asignado.",
                     "Al cerrar, se aplican en cascada las sanciones (amarillas acumuladas, rojas, dobles amarillas)."],
             aceptacion=["El cierre se bloquea e informa cuántos goles sin goleador hay por equipo.",
                         "Tras el cierre, las suspensiones correspondientes quedan registradas."]),
        dict(id="RF-ACTA-06", nombre="Bloqueo y documento del acta cerrada",
             desc="Una vez cerrada, el acta no admite ediciones y queda disponible como documento (PDF) oficial.",
             actores="Administrador de liga, Delegado",
             reglas=["El acta cerrada es de solo lectura.",
                     "Se genera el documento PDF del acta."],
             aceptacion=["No se pueden registrar ni editar incidencias en un acta cerrada.",
                         "El PDF del acta refleja el resultado e incidencias finales."]),
        dict(id="RF-ACTA-07", nombre="Acta sin conexión (offline)",
             desc="Permite cargar el acta en la cancha sin señal; la información se sincroniza al recuperar conexión.",
             actores="Personal de planilla",
             reglas=["Las acciones realizadas offline se encolan localmente.",
                     "Al volver la conexión, la cola se sincroniza con el servidor."],
             aceptacion=["El planillero opera el acta sin conexión y, al reconectar, los datos quedan persistidos."]),
     ]),

    ("3.9 Match Center en vivo",
     "Seguimiento del partido en tiempo real para operadores y público.",
     [
        dict(id="RF-MC-01", nombre="Cronómetro del partido",
             desc="Controla el tiempo del partido por periodos según la duración configurada en el torneo.",
             actores="Personal de planilla, Administrador de liga",
             reglas=["El cronómetro contempla periodos y entretiempo.",
                     "El tiempo se mantiene consistente aunque se recargue la vista."],
             aceptacion=["El cronómetro refleja el tiempo real del partido y persiste entre recargas."]),
        dict(id="RF-MC-02", nombre="Marcador en vivo",
             desc="Difunde el marcador y el estado del partido en tiempo real mediante conexión persistente (WebSocket).",
             actores="Sistema, Hincha, Operadores",
             reglas=["El marcador se actualiza en vivo al registrarse goles.",
                     "Las vistas conectadas reciben el estado actualizado sin recargar."],
             aceptacion=["Al registrar un gol, todas las pantallas conectadas muestran el nuevo marcador."]),
        dict(id="RF-MC-03", nombre="Registro rápido de goles",
             desc="Permite registrar goles ágilmente desde el panel del Match Center durante el partido.",
             actores="Personal de planilla, Administrador de liga",
             reglas=["El registro respeta la regla de goleador obligatorio (RF-ACTA-02).",
                     "El registro actualiza el marcador en vivo."],
             aceptacion=["Un gol registrado desde el panel se refleja en el marcador y en el acta."]),
        dict(id="RF-MC-04", nombre="Vista pública embebible",
             desc="Ofrece una vista pública del partido en vivo que puede incrustarse en sitios externos.",
             actores="Hincha",
             reglas=["La vista pública no requiere autenticación.",
                     "Muestra marcador y estado del partido en vivo."],
             aceptacion=["El público accede al marcador en vivo sin iniciar sesión."]),
        dict(id="RF-MC-05", nombre="Finalización del partido",
             desc="Permite marcar el partido como finalizado desde el Match Center.",
             actores="Personal de planilla, Administrador de liga",
             reglas=["La finalización del partido es un paso previo al cierre del acta.",
                     "La acción está rotulada como 'Finalizar partido'."],
             aceptacion=["El operador finaliza el partido y queda habilitado el cierre del acta."]),
     ]),

    ("3.10 Tribunal y disciplina",
     "Gestión automática y manual de sanciones disciplinarias.",
     [
        dict(id="RF-DIS-01", nombre="Sanciones automáticas",
             desc="Aplica automáticamente las sanciones derivadas de las incidencias al cerrar el acta.",
             actores="Sistema",
             reglas=["Se contemplan tarjetas rojas, dobles amarillas y acumulación de amarillas.",
                     "Las sanciones se generan en el cierre del acta (RF-ACTA-05)."],
             aceptacion=["Las sanciones correspondientes quedan registradas tras el cierre del acta."]),
        dict(id="RF-DIS-02", nombre="Suspensiones por acumulación de amarillas",
             desc="Genera una suspensión cuando un jugador alcanza el umbral de amarillas configurado en el torneo.",
             actores="Sistema",
             reglas=["El umbral es el definido en RF-TOR-05.",
                     "La suspensión inhabilita al jugador la cantidad de fechas correspondiente."],
             aceptacion=["Al alcanzar el umbral, el jugador queda suspendido automáticamente."]),
        dict(id="RF-DIS-03", nombre="Gestión de suspensiones",
             desc="Mantiene el estado de las suspensiones vigentes y su cumplimiento a lo largo de las fechas.",
             actores="Administrador de liga",
             reglas=["Las suspensiones se descuentan a medida que transcurren las fechas.",
                     "Un jugador suspendido no puede ser incluido mientras dure la sanción."],
             aceptacion=["El sistema refleja correctamente las suspensiones vigentes y cumplidas."]),
        dict(id="RF-DIS-04", nombre="Jugadores vetados (disciplina)",
             desc="Permite vetar jugadores por razones disciplinarias, integrándose con RF-JUG-06.",
             actores="Administrador de liga",
             reglas=["El veto bloquea la participación del jugador.",
                     "Queda registrado para trazabilidad."],
             aceptacion=["Un jugador vetado no puede participar hasta que se levante el veto."]),
     ]),

    ("3.11 Designaciones de personal",
     "Asignación de árbitros y personal de cancha a los partidos y jornadas.",
     [
        dict(id="RF-DES-01", nombre="Catálogo de personal",
             desc="Mantiene el registro de árbitros, planilleros, paramédicos y personal de seguridad, con sus tarifas y datos.",
             actores="Administrador de liga",
             reglas=["Cada persona tiene un rol y una tarifa asociada.",
                     "El catálogo alimenta las designaciones y los pagos."],
             aceptacion=["Es posible administrar el personal y sus tarifas."]),
        dict(id="RF-DES-02", nombre="Designación de árbitros por partido",
             desc="Asigna árbitros a cada partido según los cupos por rol.",
             actores="Administrador de liga",
             reglas=["Cada partido admite una cantidad de árbitros según su rol.",
                     "El sistema evita asignaciones incompatibles (ver RF-DES-06)."],
             aceptacion=["Cada partido queda con sus árbitros designados."]),
        dict(id="RF-DES-03", nombre="Cobertura de recinto por jornada",
             desc="Designa al personal que cubre el recinto durante toda la jornada (paramédicos y personal adicional), considerando el día completo y no cada partido.",
             actores="Administrador de liga",
             reglas=["La cobertura se asigna por jornada/día, no por partido individual.",
                     "Cuando hay dos o más torneos el mismo día, una misma persona puede cubrir todos los torneos de ese día.",
                     "El pago de la cobertura se realiza una sola vez por día aunque cubra varios torneos."],
             aceptacion=["La cobertura de recinto se designa a nivel de día completo.",
                         "Una persona compartida entre torneos del mismo día se paga una única vez."]),
        dict(id="RF-DES-04", nombre="Auto-asignación de designaciones",
             desc="Asigna automáticamente árbitros y cobertura de recinto a una jornada, equilibrando la carga.",
             actores="Administrador de liga",
             reglas=["La auto-asignación distribuye la carga entre el personal disponible.",
                     "Respeta las cantidades configuradas por jornada (RF-TOR-06) y la disponibilidad."],
             aceptacion=["La auto-asignación cubre la jornada respetando disponibilidad y cantidades configuradas."]),
        dict(id="RF-DES-05", nombre="Disponibilidad y ausencias del personal",
             desc="Registra la disponibilidad y las ausencias del personal para no designarlo cuando no puede trabajar.",
             actores="Administrador de liga, Personal",
             reglas=["El personal ausente en un día no es designado ese día.",
                     "La disponibilidad condiciona la auto-asignación."],
             aceptacion=["Una persona marcada ausente no recibe designaciones en ese día."]),
        dict(id="RF-DES-06", nombre="Detección de doble reserva",
             desc="Evita que una misma persona quede designada en dos partidos que se solapan.",
             actores="Sistema",
             reglas=["El sistema detecta solapamientos horarios al designar.",
                     "Advierte o impide la doble reserva."],
             aceptacion=["No se concreta una designación que genere un solapamiento horario."]),
        dict(id="RF-DES-07", nombre="Confirmación de asistencia",
             desc="Permite al personal confirmar su asistencia a las designaciones recibidas.",
             actores="Árbitro / Personal",
             reglas=["El personal visualiza sus designaciones y confirma asistencia.",
                     "El estado de asistencia alimenta el cálculo de pagos."],
             aceptacion=["El personal confirma su asistencia y el estado queda registrado."]),
     ]),

    ("3.12 Pagos a personal",
     "Cálculo y liquidación de los pagos al personal de la liga.",
     [
        dict(id="RF-PAGP-01", nombre="Cuentas por pagar al personal",
             desc="Genera los montos a pagar al personal en función de sus designaciones efectivamente cumplidas.",
             actores="Sistema, Administrador de liga",
             reglas=["El monto se devenga por designación con asistencia.",
                     "La cobertura de recinto compartida se paga una vez por día (RF-DES-03)."],
             aceptacion=["El sistema calcula correctamente lo adeudado a cada persona."]),
        dict(id="RF-PAGP-02", nombre="Liquidaciones",
             desc="Permite emitir y revertir liquidaciones de pago al personal.",
             actores="Administrador de liga",
             reglas=["Una liquidación consolida pagos pendientes.",
                     "Una liquidación puede revertirse, devolviendo los montos a pendientes."],
             aceptacion=["Emitir una liquidación marca los pagos como liquidados; revertirla los restituye a pendientes."]),
        dict(id="RF-PAGP-03", nombre="Datos bancarios y exportación",
             desc="Administra los datos bancarios del personal y exporta la información de pago (por ejemplo, a Excel).",
             actores="Administrador de liga",
             reglas=["Cada persona puede tener sus datos bancarios.",
                     "La exportación facilita el pago por transferencia."],
             aceptacion=["Es posible exportar la nómina de pagos con los datos necesarios."]),
        dict(id="RF-PAGP-04", nombre="Portal de pagos del personal",
             desc="Permite a cada persona consultar sus pagos recibidos y pendientes desde su portal.",
             actores="Árbitro / Personal",
             reglas=["El portal muestra el total recibido y el total pendiente.",
                     "Detalla los pagos por designación."],
             aceptacion=["El personal visualiza claramente sus pagos recibidos y pendientes."]),
     ]),

    ("3.13 Finanzas y cobros",
     "Cobros a clubes y jugadores, y conciliación de los ingresos de la liga.",
     [
        dict(id="RF-FIN-01", nombre="Tarifario del torneo",
             desc="Define los conceptos a cobrar por torneo: matrícula y cuotas.",
             actores="Administrador de liga",
             reglas=["El tarifario contempla matrícula y cuotas.",
                     "Los cobros se generan a partir del tarifario."],
             aceptacion=["Los cobros reflejan los conceptos y montos del tarifario."]),
        dict(id="RF-FIN-02", nombre="Generación de cobros",
             desc="Genera los cobros correspondientes a clubes/jugadores según el tarifario.",
             actores="Administrador de liga",
             reglas=["Cada cobro tiene un estado (pendiente, pagado).",
                     "Los cobros se asocian al club/jugador correspondiente."],
             aceptacion=["Los cobros generados quedan asociados y con su estado correcto."]),
        dict(id="RF-FIN-03", nombre="Cobranza automática (recordatorios)",
             desc="Envía recordatorios automáticos de cobranza sobre los montos pendientes.",
             actores="Sistema",
             reglas=["Los recordatorios se envían según la mora.",
                     "El proceso no afecta el aislamiento entre ligas."],
             aceptacion=["Los deudores reciben recordatorios de sus pagos pendientes."]),
        dict(id="RF-FIN-04", nombre="Pasarela de pago en línea",
             desc="Integra una pasarela de pago (por ejemplo, Flow) para que clubes y jugadores paguen en línea, con confirmación por webhook.",
             actores="Delegado, Jugador, Sistema",
             reglas=["Cada liga utiliza su propia cuenta de la pasarela.",
                     "La confirmación de pago se recibe y valida mediante webhook firmado."],
             aceptacion=["Un pago en línea aprobado actualiza el estado del cobro automáticamente."]),
        dict(id="RF-FIN-05", nombre="Boletas electrónicas (SII)",
             desc="Emite el documento tributario correspondiente al pago según la normativa chilena (SII).",
             actores="Sistema, Administrador de liga",
             reglas=["La emisión cumple los requisitos del SII.",
                     "El documento queda asociado al pago."],
             aceptacion=["Un pago genera su boleta electrónica asociada."]),
        dict(id="RF-FIN-06", nombre="Registro y conciliación de transacciones",
             desc="Registra las transacciones financieras y permite conciliar ingresos.",
             actores="Administrador de liga",
             reglas=["Cada transacción queda registrada con su origen y estado.",
                     "La conciliación relaciona pagos y cobros."],
             aceptacion=["Las transacciones son trazables y conciliables."]),
     ]),

    ("3.14 Portal público y comunidad",
     "Vistas abiertas para los hinchas y espacios de patrocinio.",
     [
        dict(id="RF-PUB-01", nombre="Fixture y resultados públicos",
             desc="Publica el calendario y los resultados de la liga para consulta abierta.",
             actores="Hincha",
             reglas=["No requiere autenticación.",
                     "Se actualiza con la operación de la liga."],
             aceptacion=["El público consulta el fixture y los resultados sin iniciar sesión."]),
        dict(id="RF-PUB-02", nombre="Tabla de posiciones",
             desc="Muestra la tabla de posiciones por categoría/serie con los criterios de desempate configurados.",
             actores="Hincha",
             reglas=["La tabla respeta los tiebreakers del torneo (RF-TOR-03).",
                     "Se presenta por categoría/serie."],
             aceptacion=["La tabla pública refleja el orden correcto según los criterios definidos."]),
        dict(id="RF-PUB-03", nombre="Rankings y estadísticas",
             desc="Presenta rankings (por ejemplo, goleadores) y estadísticas de la competencia.",
             actores="Hincha",
             reglas=["Los rankings se derivan de las incidencias registradas.",
                     "Las estadísticas se actualizan con los partidos."],
             aceptacion=["Los rankings y estadísticas reflejan los datos de los partidos."]),
        dict(id="RF-PUB-04", nombre="Sponsors y banners",
             desc="Permite mostrar espacios publicitarios de patrocinadores dentro de la liga.",
             actores="Administrador de liga",
             reglas=["Los banners se administran por liga.",
                     "Constituyen un espacio monetizable."],
             aceptacion=["Los banners configurados se muestran en los espacios definidos."]),
     ]),

    ("3.15 Notificaciones",
     "Comunicación con los usuarios por distintos canales.",
     [
        dict(id="RF-NOT-01", nombre="Correos transaccionales",
             desc="Envía correos para invitaciones, recuperación de contraseña y avisos relevantes.",
             actores="Sistema",
             reglas=["Los correos se envían mediante un proveedor transaccional.",
                     "Cada correo se ajusta al idioma del usuario."],
             aceptacion=["Los eventos relevantes generan el correo correspondiente."]),
        dict(id="RF-NOT-02", nombre="Notificaciones push",
             desc="Envía notificaciones push a los dispositivos de los usuarios para avisos en tiempo oportuno.",
             actores="Sistema",
             reglas=["El usuario puede recibir notificaciones en su dispositivo.",
                     "Se respetan los permisos de notificación."],
             aceptacion=["El usuario suscrito recibe las notificaciones push pertinentes."]),
     ]),

    ("3.16 Reportería e informes",
     "Información consolidada para la toma de decisiones.",
     [
        dict(id="RF-REP-01", nombre="Informes de la operación",
             desc="Genera informes de disciplina, finanzas, competición y arbitraje/personal.",
             actores="Administrador de liga",
             reglas=["Cada informe consolida la información de su área.",
                     "Los informes consideran el alcance de la liga."],
             aceptacion=["Los informes presentan datos consistentes con la operación."]),
        dict(id="RF-REP-02", nombre="Exportación a Excel y PDF",
             desc="Permite exportar los informes a formatos Excel y PDF para su distribución.",
             actores="Administrador de liga",
             reglas=["La exportación conserva la estructura del informe.",
                     "Disponible para los informes principales."],
             aceptacion=["Los informes se exportan correctamente a Excel y PDF."]),
     ]),

    ("3.17 Auditoría",
     "Trazabilidad de las acciones críticas del sistema.",
     [
        dict(id="RF-AUD-01", nombre="Registro de auditoría",
             desc="Registra las acciones críticas (creación/edición/eliminación de entidades de dominio, inicios de sesión, cambios de rol, pagos, cierre de actas) con su contexto.",
             actores="Sistema",
             reglas=["Cada registro incluye usuario, liga, acción, marca de tiempo e IP.",
                     "Las sesiones de impersonación quedan auditadas."],
             aceptacion=["Las acciones críticas dejan un registro de auditoría consultable."]),
     ]),
]

for titulo, intro, reqs in MODULOS:
    doc.add_heading(titulo, level=2)
    doc.add_paragraph(intro)
    for rf in reqs:
        render_req(rf)

doc.add_page_break()

# ================================================================== 4. RNF
doc.add_heading('4. Requerimientos no funcionales', level=1)
RNF = [
    ("RNF-SEG", "Seguridad y aislamiento multi-tenant",
     "El sistema protege los datos de cada liga y de sus usuarios con controles de seguridad de nivel productivo.",
     ["Aislamiento de datos por liga a nivel del motor de base de datos (Row-Level Security), de modo que una liga nunca accede a datos de otra.",
      "Autenticación con tokens de acceso de corta vida y refresco rotativo; contraseñas con hash robusto (bcrypt).",
      "Límite de intentos de inicio de sesión, cabeceras de seguridad HTTP y lista blanca de orígenes (CORS).",
      "Webhooks de pago verificados por firma.",
      "Registro de auditoría de acciones críticas e impersonación de soporte auditada.",
      "Cumplimiento de la Ley N° 19.628 sobre protección de la vida privada (datos personales, Chile)."]),
    ("RNF-DISP", "Disponibilidad y operación",
     "El sistema está diseñado para operar de forma continua y recuperarse ante fallos.",
     ["Verificaciones de salud (liveness/readiness) y de versión para soportar despliegues controlados.",
      "Reinicio automático de servicios y arranque ordenado.",
      "Respaldo periódico de la base de datos con política de retención."]),
    ("RNF-OBS", "Observabilidad",
     "La operación es monitoreable para detectar y diagnosticar problemas.",
     ["Registro estructurado de eventos con identificadores de solicitud, usuario y liga.",
      "Captura centralizada de errores y métricas de rendimiento."]),
    ("RNF-PER", "Rendimiento",
     "El sistema responde con fluidez en los flujos de uso habituales.",
     ["Las vistas presentan estados de carga (esqueletos) y evitan pantallas en blanco.",
      "El marcador y el Match Center se actualizan en tiempo real."]),
    ("RNF-USA", "Usabilidad y diseño móvil",
     "La experiencia está pensada para uso en cancha (móvil) y en oficina (escritorio).",
     ["Diseño centrado en móvil, con objetivos táctiles adecuados.",
      "Aplicación instalable (PWA) con capacidad de uso sin conexión para el acta."]),
    ("RNF-I18N", "Internacionalización",
     "El sistema soporta múltiples idiomas para habilitar su expansión regional.",
     ["Idiomas soportados: español, inglés y portugués.",
      "Las comunicaciones se ajustan al idioma del usuario."]),
    ("RNF-MAN", "Mantenibilidad",
     "El sistema es sostenible en el tiempo y seguro de evolucionar.",
     ["Cambios de esquema gestionados mediante migraciones versionadas; sin modificaciones manuales en producción.",
      "Tipos compartidos y reglas de dominio centralizadas para consistencia entre cliente y servidor."]),
]
for code, nombre, desc, items in RNF:
    doc.add_heading(f"{code} — {nombre}", level=3)
    doc.add_paragraph(desc)
    for b in items:
        doc.add_paragraph(b, style='List Bullet')

doc.add_page_break()

# ================================================================== 5. ROLES
doc.add_heading('5. Modelo de roles y permisos', level=1)
doc.add_paragraph('La siguiente matriz resume las capacidades principales por rol (✓ = capacidad habilitada).')
roles = ['Super Admin', 'Admin liga', 'Delegado', 'Personal', 'Jugador', 'Hincha']
caps = [
    ('Administrar ligas y planes', [1, 0, 0, 0, 0, 0]),
    ('Configurar la liga', [1, 1, 0, 0, 0, 0]),
    ('Crear torneos y fixture', [1, 1, 0, 0, 0, 0]),
    ('Designar personal', [1, 1, 0, 0, 0, 0]),
    ('Cargar y cerrar actas', [1, 1, 0, 1, 0, 0]),
    ('Gestionar cobros', [1, 1, 0, 0, 0, 0]),
    ('Liquidar pagos a personal', [1, 1, 0, 0, 0, 0]),
    ('Gestionar plantel del club', [1, 1, 1, 0, 0, 0]),
    ('Consultar designaciones y pagos propios', [0, 0, 0, 1, 0, 0]),
    ('Consultar ficha y calendario propios', [0, 0, 0, 0, 1, 0]),
    ('Ver portal público', [1, 1, 1, 1, 1, 1]),
]
tbl = doc.add_table(rows=1 + len(caps), cols=1 + len(roles))
tbl.style = 'Light Grid Accent 1'
set_cell_text(tbl.rows[0].cells[0], 'Capacidad', bold=True, white=True, size=9)
shade(tbl.rows[0].cells[0], '103D2B')
for j, rn in enumerate(roles):
    set_cell_text(tbl.rows[0].cells[1 + j], rn, bold=True, white=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(tbl.rows[0].cells[1 + j], '103D2B')
for i, (cap, vals) in enumerate(caps):
    set_cell_text(tbl.rows[1 + i].cells[0], cap, size=9)
    for j, v in enumerate(vals):
        set_cell_text(tbl.rows[1 + i].cells[1 + j], '✓' if v else '—', size=10,
                      align=WD_ALIGN_PARAGRAPH.CENTER, color=VERDE2 if v else GRIS, bold=bool(v))
tbl.columns[0].width = Cm(6.2)

doc.add_page_break()

# ================================================================== 6. GLOSARIO
doc.add_heading('6. Glosario', level=1)
gloss = [
    ('Tenant (liga)', 'Inquilino del sistema; cada liga es un tenant con datos aislados.'),
    ('Multi-tenant', 'Arquitectura en la que múltiples ligas comparten la plataforma manteniendo sus datos separados.'),
    ('RLS (Row-Level Security)', 'Mecanismo del motor de base de datos que aísla las filas por liga.'),
    ('Fixture', 'Calendario de partidos del torneo.'),
    ('Fecha / jornada', 'Conjunto de partidos programados en un mismo periodo del torneo.'),
    ('Acta', 'Registro oficial del desarrollo y resultado de un partido.'),
    ('Planilla', 'Lista de jugadores habilitados de un equipo para un partido.'),
    ('Incidencia', 'Hecho del partido registrado en el acta (gol, autogol, amarilla, roja).'),
    ('Walkover', 'Resultado reglamentario por no presentación de un equipo.'),
    ('Designación', 'Asignación de un árbitro o personal a un partido o jornada.'),
    ('Cobertura de recinto', 'Personal asignado al recinto por jornada (paramédicos, seguridad y otros).'),
    ('Liquidación', 'Consolidación de pagos pendientes del personal para su pago.'),
    ('Dunning / cobranza', 'Proceso de recordatorios de pagos pendientes.'),
    ('Pasarela de pago', 'Servicio externo que procesa los pagos en línea.'),
    ('Boleta SII', 'Documento tributario electrónico exigido por el Servicio de Impuestos Internos (Chile).'),
    ('PWA', 'Aplicación web instalable con capacidades sin conexión.'),
    ('JWT', 'Token de autenticación firmado utilizado para sostener la sesión.'),
    ('Audit log', 'Registro de auditoría de acciones críticas.'),
]
gt = doc.add_table(rows=1 + len(gloss), cols=2)
gt.style = 'Light List Accent 1'
set_cell_text(gt.rows[0].cells[0], 'Término', bold=True, white=True, size=10)
set_cell_text(gt.rows[0].cells[1], 'Definición', bold=True, white=True, size=10)
shade(gt.rows[0].cells[0], '103D2B'); shade(gt.rows[0].cells[1], '103D2B')
for i, (t, d) in enumerate(gloss):
    set_cell_text(gt.rows[1 + i].cells[0], t, bold=True, color=VERDE, size=10)
    set_cell_text(gt.rows[1 + i].cells[1], d, size=10)
gt.columns[0].width = Cm(4.6); gt.columns[1].width = Cm(11.4)


# ------------------------------------------------------------------ footer
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.text = 'LigaPlus — Especificación de Requerimientos Funcionales · Confidencial'
fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = GRIS
fp.add_run('\t\tPágina ')
fp.runs[1].font.size = Pt(8); fp.runs[1].font.color.rgb = GRIS
add_field(fp, 'PAGE', '1')

enable_update_fields()

OUT = r'C:\Claude\Dev\Fixtura\docs\LigaPlus-Requerimientos-Funcionales.docx'
doc.save(OUT)
n_rf = sum(len(m[2]) for m in MODULOS)
print('OK')
print('Archivo:', OUT)
print('Modulos funcionales:', len(MODULOS))
print('Requerimientos funcionales:', n_rf)
print('Requerimientos no funcionales:', len(RNF))
